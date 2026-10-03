"""Reading source documents, term bases and style sheets; writing deliverables."""

import csv
import io
import json
import os
import re
from typing import Dict, List

MAX_FILE_SIZE = int(os.getenv("MAX_FILE_SIZE", str(25 * 1024 * 1024)))

DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".json", ".html", ".htm", ".txt", ".md"}
GLOSSARY_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".json", ".txt"}
STYLE_SHEET_EXTENSIONS = {".txt", ".md", ".docx", ".pdf", ".html", ".htm", ".json"}
OUTPUT_FORMATS = {"docx", "txt", "md", "json", "html"}


class FileError(ValueError):
    pass


def _ext(path: str) -> str:
    return os.path.splitext(path)[1].lower()


def validate(path: str, allowed: set, label: str):
    if os.path.getsize(path) > MAX_FILE_SIZE:
        raise FileError(f"{label} exceeds the {MAX_FILE_SIZE // (1024 * 1024)} MB limit.")
    if _ext(path) not in allowed:
        raise FileError(f"Unsupported {label.lower()} type '{_ext(path)}'. Allowed: {', '.join(sorted(allowed))}")


def extract_text(path: str) -> str:
    ext = _ext(path)
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
    if ext == ".docx":
        from docx import Document

        return "\n\n".join(p.text for p in Document(path).paragraphs if p.text.strip())
    if ext == ".pptx":
        from pptx import Presentation

        blocks = []
        for slide in Presentation(path).slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    blocks.append(shape.text_frame.text.strip())
        return "\n\n".join(blocks)
    if ext in (".html", ".htm"):
        from bs4 import BeautifulSoup

        with open(path, "r", encoding="utf-8", errors="replace") as f:
            soup = BeautifulSoup(f.read(), "html.parser")
        return soup.get_text("\n\n", strip=True)
    if ext == ".json":
        with open(path, "r", encoding="utf-8") as f:
            return json.dumps(json.load(f), ensure_ascii=False, indent=2)
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


_SEPARATORS = ["\n\n", "\n", r"(?<=[.!?。！？])\s+", " "]


def split_text(text: str, chunk_size: int) -> List[str]:
    """Recursive splitter that keeps paragraphs and sentences whole where possible.

    No overlap: overlapping chunks would duplicate text in the joined translation.
    """
    text = text.strip()
    if not text:
        return []
    chunk_size = max(200, chunk_size)
    return [c.strip() for c in _split(text, chunk_size, 0) if c.strip()]


def _split(text: str, size: int, level: int) -> List[str]:
    if len(text) <= size:
        return [text]
    if level >= len(_SEPARATORS):
        return [text[i : i + size] for i in range(0, len(text), size)]
    sep = _SEPARATORS[level]
    joiner = sep if not sep.startswith("(") else " "
    pieces = re.split(sep, text) if sep.startswith("(") else text.split(sep)
    chunks, current = [], ""
    for piece in pieces:
        candidate = f"{current}{joiner}{piece}" if current else piece
        if len(candidate) <= size:
            current = candidate
            continue
        if current:
            chunks.append(current)
        if len(piece) > size:
            chunks.extend(_split(piece, size, level + 1))
            current = ""
        else:
            current = piece
    if current:
        chunks.append(current)
    return chunks


def parse_glossary(path: str) -> Dict[str, str]:
    """First column = source term, second column = target term. A header row is tolerated."""
    ext = _ext(path)
    rows: List[List[str]] = []
    if ext == ".xlsx":
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = [[("" if c is None else str(c)) for c in r] for r in sheet.iter_rows(values_only=True)]
    elif ext == ".json":
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            rows = [[k, v] for k, v in data.items()]
        elif isinstance(data, list):
            for item in data:
                if isinstance(item, dict):
                    vals = list(item.values())
                    if len(vals) >= 2:
                        rows.append([vals[0], vals[1]])
                elif isinstance(item, (list, tuple)) and len(item) >= 2:
                    rows.append([item[0], item[1]])
    else:
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            content = f.read()
        if ext == ".txt" and not any(d in content for d in ("\t", ",", ";")):
            rows = [re.split(r"\s*(?:=|→|->|:)\s*", line, maxsplit=1) for line in content.splitlines()]
        else:
            try:
                dialect = csv.Sniffer().sniff(content[:2048], delimiters=",\t;|")
            except csv.Error:
                dialect = csv.excel_tab if ext == ".tsv" else csv.excel
            rows = list(csv.reader(io.StringIO(content), dialect))

    glossary: Dict[str, str] = {}
    for row in rows:
        if len(row) < 2:
            continue
        src, tgt = str(row[0]).strip(), str(row[1]).strip()
        if src and tgt:
            glossary[src] = tgt
    if glossary:
        first_key = next(iter(glossary))
        if first_key.lower() in {"source", "term", "source term", "english", "src"}:
            glossary.pop(first_key)
    if not glossary:
        raise FileError("The terminology list is empty or could not be read (need two columns: source, target).")
    return glossary


def relevant_terms(glossary: Dict[str, str], text: str) -> Dict[str, str]:
    lowered = text.lower()
    return {k: v for k, v in glossary.items() if k.lower() in lowered}


def save_translation(text_blocks: List[str], out_dir: str, stem: str, fmt: str, meta: dict) -> str:
    fmt = fmt.lower() if fmt.lower() in OUTPUT_FORMATS else "docx"
    path = os.path.join(out_dir, f"{stem}_translated.{fmt}")
    joined = "\n\n".join(text_blocks)
    if fmt == "docx":
        from docx import Document

        doc = Document()
        for block in text_blocks:
            for para in block.split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
        doc.core_properties.title = f"{stem} ({meta.get('target_lang', '')})"
        doc.core_properties.author = "Agentic AI Translation Company"
        doc.save(path)
    elif fmt == "json":
        with open(path, "w", encoding="utf-8") as f:
            json.dump({**meta, "translation": joined, "chunks": text_blocks}, f, ensure_ascii=False, indent=2)
    elif fmt == "html":
        import html

        paras = "".join(f"<p>{html.escape(p.strip())}</p>\n" for p in joined.split("\n\n") if p.strip())
        with open(path, "w", encoding="utf-8") as f:
            f.write(
                f'<!doctype html>\n<html lang="en"><head><meta charset="utf-8"><title>{html.escape(stem)}</title></head>'
                f"<body>\n{paras}</body></html>\n"
            )
    else:
        with open(path, "w", encoding="utf-8") as f:
            f.write(joined + "\n")
    return path


def save_details(details: dict, out_dir: str, stem: str) -> str:
    path = os.path.join(out_dir, f"{stem}_details.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(details, f, ensure_ascii=False, indent=2)
    return path
